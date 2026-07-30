"""
Document management API — upload, list, delete resumes/JDs/account notes.

Upload flow:
  1. Client uploads a file (PDF, DOCX, TXT) with metadata.
  2. Server extracts plain text.
  3. Text is chunked (overlapping windows).
  4. Each chunk is embedded via the Embedder provider.
  5. Chunks + embeddings are stored in pgvector via the VectorStore.

This is the ingestion side of the RAG pipeline.
"""

from __future__ import annotations

import io
import logging
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel

from app.auth.middleware import AuthUser, get_optional_user
import app.db.engine as db_engine
from app.db.models import Document, DocumentChunk
import app.providers as providers
from app.providers.vector_store import VectorChunk

logger = logging.getLogger(__name__)
router = APIRouter()

# Chunk config
CHUNK_SIZE = 500  # characters
CHUNK_OVERLAP = 100


class DocumentResponse(BaseModel):
    id: str
    doc_type: str
    filename: str
    label: str | None
    chunk_count: int


class DocumentListResponse(BaseModel):
    documents: list[DocumentResponse]


def _chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        if chunk.strip():
            chunks.append(chunk.strip())
        start = end - overlap
    return chunks


async def _extract_text(file: UploadFile) -> str:
    """Extract plain text from uploaded file."""
    content = await file.read()
    filename = file.filename or ""

    if filename.endswith(".txt"):
        return content.decode("utf-8", errors="replace")
    elif filename.endswith(".pdf"):
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            # Fallback: try to decode as text
            return content.decode("utf-8", errors="replace")
    elif filename.endswith((".docx", ".doc")):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            return content.decode("utf-8", errors="replace")
    else:
        # Best effort: try text decoding
        return content.decode("utf-8", errors="replace")


@router.post("", response_model=DocumentResponse)
async def upload_document(
    file: UploadFile = File(...),
    doc_type: str = Form(...),
    label: str | None = Form(None),
    user: AuthUser | None = Depends(get_optional_user),
) -> DocumentResponse:
    """Upload and ingest a document for RAG grounding."""
    if doc_type not in ("resume", "jd", "account_note", "other"):
        raise HTTPException(400, "Invalid doc_type")

    # Extract text
    raw_text = await _extract_text(file)
    if not raw_text.strip():
        raise HTTPException(400, "Could not extract text from file")

    # Chunk
    chunks = _chunk_text(raw_text)
    doc_id = str(uuid.uuid4())
    user_id = str(user.id) if user else "00000000-0000-0000-0000-000000000000"

    # Embed all chunks (graceful fallback if embedder unavailable)
    try:
        embeddings = await providers.embedder.embed(chunks)
    except Exception as e:
        logger.warning("Embedding failed, using zero vectors: %s", e)
        dim = 384
        embeddings = [[0.0] * dim for _ in chunks]

    # Store document metadata
    if db_engine.async_session_factory is None:
        raise HTTPException(503, "Database not initialized")
    async with db_engine.async_session_factory() as session:
        doc = Document(
            id=uuid.UUID(doc_id),
            user_id=uuid.UUID(user_id),
            doc_type=doc_type,
            filename=file.filename or "unknown",
            label=label,
            raw_text=raw_text,
            chunk_count=len(chunks),
        )
        session.add(doc)
        await session.commit()

    # Store chunks in vector store
    vector_chunks = [
        VectorChunk(
            id=str(uuid.uuid4()),
            document_id=doc_id,
            text=chunk_text,
            embedding=embeddings[i],
            metadata={"chunk_index": str(i), "doc_type": doc_type},
        )
        for i, chunk_text in enumerate(chunks)
    ]
    try:
        await providers.vector_store.upsert(vector_chunks)
    except Exception as e:
        logger.warning("Vector store upsert failed (doc still saved): %s", e)

    logger.info("Ingested document %s: %d chunks from %s", doc_id, len(chunks), file.filename)

    return DocumentResponse(
        id=doc_id,
        doc_type=doc_type,
        filename=file.filename or "unknown",
        label=label,
        chunk_count=len(chunks),
    )


@router.get("", response_model=DocumentListResponse)
async def list_documents(
    user: AuthUser | None = Depends(get_optional_user),
) -> DocumentListResponse:
    """List all documents for a user."""
    user_id = str(user.id) if user else "00000000-0000-0000-0000-000000000000"
    assert db_engine.async_session_factory is not None
    async with db_engine.async_session_factory() as session:
        from sqlalchemy import select
        result = await session.execute(
            select(Document).where(Document.user_id == uuid.UUID(user_id))
        )
        docs = result.scalars().all()

    return DocumentListResponse(
        documents=[
            DocumentResponse(
                id=str(d.id),
                doc_type=d.doc_type,
                filename=d.filename,
                label=d.label,
                chunk_count=d.chunk_count,
            )
            for d in docs
        ]
    )


@router.delete("/{document_id}")
async def delete_document(document_id: str) -> dict[str, str]:
    """Delete a document and its chunks."""
    # Delete vector chunks
    deleted = await providers.vector_store.delete_by_document(document_id)

    # Delete document record
    assert db_engine.async_session_factory is not None
    async with db_engine.async_session_factory() as session:
        from sqlalchemy import delete
        await session.execute(
            delete(Document).where(Document.id == uuid.UUID(document_id))
        )
        await session.commit()

    return {"status": "deleted", "chunks_removed": str(deleted)}
